from fastapi import FastAPI, APIRouter, UploadFile, File, HTTPException
from fastapi.responses import StreamingResponse
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel, ConfigDict
from dotenv import load_dotenv
from sqlalchemy import create_engine, String, Text, Integer, select, func, or_
from sqlalchemy.orm import DeclarativeBase, Mapped, mapped_column, sessionmaker
from pathlib import Path
import mimetypes
import csv, io, os, uuid

ROOT = Path(__file__).parent
load_dotenv(ROOT / ".env")
UPLOADS = ROOT / "uploads"
UPLOADS.mkdir(exist_ok=True)
DATABASE_URL = os.environ["DATABASE_URL"]
engine = create_engine(DATABASE_URL, connect_args={"check_same_thread": False} if DATABASE_URL.startswith("sqlite") else {})
SessionLocal = sessionmaker(bind=engine, expire_on_commit=False)

class Base(DeclarativeBase): pass
class Collection(Base):
    __tablename__ = "collections"
    id: Mapped[int] = mapped_column(Integer, primary_key=True)
    nomor_koleksi: Mapped[str] = mapped_column(String(80), unique=True, index=True)
    nama_ilmiah: Mapped[str] = mapped_column(String(180)); nama_umum: Mapped[str] = mapped_column(String(180), default="")
    kelompok: Mapped[str] = mapped_column(String(40)); jenis_koleksi: Mapped[str] = mapped_column(String(80))
    lokasi_asal: Mapped[str] = mapped_column(String(180), default=""); kabupaten: Mapped[str] = mapped_column(String(120), default=""); provinsi: Mapped[str] = mapped_column(String(120), default=""); negara: Mapped[str] = mapped_column(String(100), default="Indonesia")
    tanggal_koleksi: Mapped[str] = mapped_column(String(20), default=""); kolektor: Mapped[str] = mapped_column(String(160), default="")
    sumber_perolehan: Mapped[str] = mapped_column(String(160), default=""); tanggal_diterima: Mapped[str] = mapped_column(String(20), default=""); keterangan_perolehan: Mapped[str] = mapped_column(Text, default="")
    kondisi_spesimen: Mapped[str] = mapped_column(String(80), default="Baik"); kelengkapan: Mapped[str] = mapped_column(String(80), default=""); kerusakan: Mapped[str] = mapped_column(String(160), default=""); catatan_kondisi: Mapped[str] = mapped_column(Text, default="")
    gedung: Mapped[str] = mapped_column(String(80), default=""); ruang: Mapped[str] = mapped_column(String(80), default=""); lemari: Mapped[str] = mapped_column(String(80), default=""); rak: Mapped[str] = mapped_column(String(80), default=""); kotak_posisi: Mapped[str] = mapped_column(String(100), default="")
    foto_path: Mapped[str] = mapped_column(String(255), default=""); keterangan_foto: Mapped[str] = mapped_column(Text, default="")
    osteologi_kelengkapan: Mapped[str] = mapped_column(String(80), default=""); bagian_tulang: Mapped[str] = mapped_column(Text, default=""); catatan_osteologi: Mapped[str] = mapped_column(Text, default="")

Base.metadata.create_all(engine)
app = FastAPI(title="Database Koleksi Zoologi UGM")
api = APIRouter(prefix="/api")
app.add_middleware(CORSMiddleware, allow_origins=["*"], allow_methods=["*"], allow_headers=["*"])
FIELDS = [c.name for c in Collection.__table__.columns if c.name != "id"]

class CollectionIn(BaseModel):
    model_config = ConfigDict(extra="ignore")
    nomor_koleksi: str; nama_ilmiah: str; nama_umum: str = ""; kelompok: str; jenis_koleksi: str
    lokasi_asal: str = ""; kabupaten: str = ""; provinsi: str = ""; negara: str = "Indonesia"; tanggal_koleksi: str = ""; kolektor: str = ""
    sumber_perolehan: str = ""; tanggal_diterima: str = ""; keterangan_perolehan: str = ""; kondisi_spesimen: str = "Baik"; kelengkapan: str = ""; kerusakan: str = ""; catatan_kondisi: str = ""
    gedung: str = ""; ruang: str = ""; lemari: str = ""; rak: str = ""; kotak_posisi: str = ""; keterangan_foto: str = ""
    osteologi_kelengkapan: str = ""; bagian_tulang: str = ""; catatan_osteologi: str = ""
def item(c): return {k: getattr(c, k) for k in FIELDS} | {"id": c.id, "foto_url": f"/api/files/{c.foto_path}" if c.foto_path else ""}

@api.get("/")
def root(): return {"message": "Museum Zoologi API aktif"}
@api.get("/collections")
def list_collections(q: str = "", kelompok: str = "", jenis: str = ""):
    with SessionLocal() as s:
        stmt = select(Collection).order_by(Collection.id.desc())
        if q: stmt = stmt.where(or_(*[getattr(Collection, f).ilike(f"%{q}%") for f in ["nomor_koleksi", "nama_ilmiah", "nama_umum", "kelompok", "jenis_koleksi", "lokasi_asal"]]))
        if kelompok: stmt = stmt.where(Collection.kelompok == kelompok)
        if jenis: stmt = stmt.where(Collection.jenis_koleksi == jenis)
        return [item(c) for c in s.scalars(stmt).all()]
@api.get("/collections/{cid}")
def get_collection(cid: int):
    with SessionLocal() as s:
        c = s.get(Collection, cid)
        if not c: raise HTTPException(404, "Koleksi tidak ditemukan")
        return item(c)
@api.post("/collections")
def create_collection(data: CollectionIn):
    with SessionLocal() as s:
        if s.scalar(select(Collection).where(Collection.nomor_koleksi == data.nomor_koleksi)): raise HTTPException(409, "Nomor koleksi sudah digunakan")
        c = Collection(**data.model_dump()); s.add(c); s.commit(); s.refresh(c); return item(c)
@api.put("/collections/{cid}")
def update_collection(cid: int, data: CollectionIn):
    with SessionLocal() as s:
        c = s.get(Collection, cid)
        if not c: raise HTTPException(404, "Koleksi tidak ditemukan")
        if s.scalar(select(Collection).where(Collection.nomor_koleksi == data.nomor_koleksi, Collection.id != cid)): raise HTTPException(409, "Nomor koleksi sudah digunakan")
        for k, v in data.model_dump().items(): setattr(c, k, v)
        s.commit(); s.refresh(c); return item(c)
@api.delete("/collections/{cid}")
def delete_collection(cid: int):
    with SessionLocal() as s:
        c = s.get(Collection, cid)
        if not c: raise HTTPException(404, "Koleksi tidak ditemukan")
        s.delete(c); s.commit(); return {"ok": True}
@api.get("/stats")
def stats():
    with SessionLocal() as s:
        groups = ["Mammalia", "Aves", "Reptilia", "Amphibia", "Pisces", "Invertebrata"]
        return {"total": s.scalar(select(func.count()).select_from(Collection)) or 0, **{g: s.scalar(select(func.count()).select_from(Collection).where(Collection.kelompok == g)) or 0 for g in groups}}
@api.post("/collections/{cid}/photo")
async def upload_photo(cid: int, file: UploadFile = File(...)):
    with SessionLocal() as s:
        c = s.get(Collection, cid)
        if not c: raise HTTPException(404, "Koleksi tidak ditemukan")
        ext = Path(file.filename or "foto.jpg").suffix.lower() or ".jpg"; name = f"{uuid.uuid4().hex}{ext}"
        (UPLOADS / name).write_bytes(await file.read()); c.foto_path = name; s.commit(); return item(c)
@api.get("/files/{name}")
def file_response(name: str):
    path = UPLOADS / Path(name).name
    if not path.exists(): raise HTTPException(404, "Foto tidak ditemukan")
    return StreamingResponse(open(path, "rb"), media_type=mimetypes.guess_type(path.name)[0] or "application/octet-stream")
@api.post("/import")
async def import_file(file: UploadFile = File(...)):
    raw = await file.read(); filename = (file.filename or "").lower()
    if filename.endswith(".csv"): rows = list(csv.DictReader(io.StringIO(raw.decode("utf-8-sig"))))
    else:
        from openpyxl import load_workbook
        values = list(load_workbook(io.BytesIO(raw), read_only=True, data_only=True).active.values); headers = [str(x or "") for x in values[0]]; rows = [dict(zip(headers, r)) for r in values[1:]]
    added, duplicates, invalid = 0, [], []
    with SessionLocal() as s:
        existing = {x[0] for x in s.execute(select(Collection.nomor_koleksi)).all()}
        for row in rows:
            r = {str(k).strip(): ("" if v is None else str(v)) for k, v in row.items()}; number = r.get("nomor_koleksi") or r.get("Nomor Koleksi")
            if not number or number in existing: duplicates.append(number or "(kosong)"); continue
            data = {k: r.get(k, "") for k in FIELDS}; data.update(nomor_koleksi=number, nama_ilmiah=r.get("nama_ilmiah", r.get("Nama Ilmiah", "")), kelompok=r.get("kelompok", r.get("Kelompok", "")), jenis_koleksi=r.get("jenis_koleksi", r.get("Jenis Koleksi", "")))
            if not data["nama_ilmiah"] or not data["kelompok"] or not data["jenis_koleksi"]: invalid.append(number or "(tanpa nomor)"); continue
            s.add(Collection(**data)); existing.add(number); added += 1
        s.commit()
    return {"added": added, "duplicates": duplicates, "invalid": invalid}
def export_rows():
    with SessionLocal() as s: return [item(c) for c in s.scalars(select(Collection).order_by(Collection.id)).all()]
@api.get("/export/csv")
def export_csv():
    out = io.StringIO(); fields = ["nomor_koleksi", "nama_ilmiah", "nama_umum", "kelompok", "jenis_koleksi", "lokasi_asal", "kondisi_spesimen", "gedung", "ruang"]; w = csv.DictWriter(out, fieldnames=fields); w.writeheader(); w.writerows([{k: r.get(k, "") for k in fields} for r in export_rows()]); return StreamingResponse(iter([out.getvalue().encode("utf-8-sig")]), media_type="text/csv", headers={"Content-Disposition": "attachment; filename=koleksi-zoologi.csv"})
@api.get("/export/xlsx")
def export_xlsx():
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.append(FIELDS)
    for r in export_rows(): ws.append([r.get(k, "") for k in FIELDS])
    out = io.BytesIO(); wb.save(out); out.seek(0); return StreamingResponse(out, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": "attachment; filename=koleksi-zoologi.xlsx"})
@api.get("/export/pdf")
def export_pdf():
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.pdfgen.canvas import Canvas
    out = io.BytesIO(); c = Canvas(out, pagesize=landscape(A4)); c.setFont("Helvetica-Bold", 14); c.drawString(36, 560, "Daftar Koleksi Zoologi Museum Biologi UGM"); c.setFont("Helvetica", 8); y = 540
    for r in export_rows():
        c.drawString(36, y, f"{r['nomor_koleksi']} | {r['nama_ilmiah']} | {r['nama_umum']} | {r['kelompok']} | {r['jenis_koleksi']}"); y -= 14
        if y < 35: c.showPage(); y = 560
    c.save(); out.seek(0); return StreamingResponse(out, media_type="application/pdf", headers={"Content-Disposition": "attachment; filename=koleksi-zoologi.pdf"})
app.include_router(api)